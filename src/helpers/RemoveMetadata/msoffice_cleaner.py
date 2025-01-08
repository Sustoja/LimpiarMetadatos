from datetime import datetime
from docx import Document
from openpyxl import load_workbook
from pathlib import Path
from pptx import Presentation

from .constants import META_FIELDS, DATE_FIELDS
from ..Logger import mylogger


def _get_document_and_property_objects(filename: Path|str):
    try:
        extension = filename.suffix.lower()
        if extension == ".xlsx":
            document = load_workbook(filename)
            return document, document.properties
        elif extension == ".docx":
            document = Document(filename)
            return document, document.core_properties
        elif extension == ".pptx":
            document = Presentation(filename)
            return document, document.core_properties
        else:
            return None, None
    except Exception as e:
        mylogger.log.error(f"Error al acceder a las propiedades de {filename}: {e}")
        return None, None

def print_msoffice_metadata(filename: Path) -> None:
    try:
        _, props = _get_document_and_property_objects(filename)
        if not props:
            mylogger.log.error(f"No se pudieron cargar las propiedades para {filename}")
            return

        extension = filename.suffix.lower()
        for field in META_FIELDS[extension]:
            mylogger.log.info(f'{field}: {getattr(props, field, None)}')
        for field in DATE_FIELDS[extension]:
            mylogger.log.info(f'{field}: {getattr(props, field, None)}')
    except Exception as e:
        mylogger.log.error(f"Error al leer los metadatos de {filename}: {e}")

def remove_msoffice_metadata(filename: Path) -> bool:
    try:
        document, props = _get_document_and_property_objects(filename)
        if not props:
            mylogger.log.error(f"No se pudieron cargar las propiedades para {filename}")
            return False

        reference_date = datetime(1900, 1, 1, 00, 00)
        extension = filename.suffix.lower()

        # Campos generales
        setattr(props, "revision", 1)
        for field in META_FIELDS.get(extension, []):
            if field != "revision":
                setattr(props, field, "")

        # Campos de fecha
        for field in DATE_FIELDS.get(extension, []):
            setattr(props, field, reference_date)

        document.save(filename)
        return True
    except Exception as e:
        mylogger.log.error(f"Error al borrar los metadatos del archivo {filename}: {e}")
        return False
